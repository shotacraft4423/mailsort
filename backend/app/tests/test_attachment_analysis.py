from __future__ import annotations

import io
import zipfile

from app.services import attachment_analysis_service as svc


def test_classify_kind_by_filename_keywords():
    assert svc.classify_kind(file_name="山田太郎_スキルシート.xlsx", extracted_text="") == "skill_sheet"
    assert svc.classify_kind(file_name="2024年8月分_請求書.pdf", extracted_text="") == "invoice"
    assert svc.classify_kind(file_name="基本契約書_20240801.docx", extracted_text="") == "contract"
    assert svc.classify_kind(file_name="random.pdf", extracted_text="今月の飲み会のお知らせです") == "other"


def test_classify_kind_falls_back_to_extracted_text_when_filename_is_generic():
    assert svc.classify_kind(file_name="document.pdf", extracted_text="本書は個別契約書として...") == "contract"


def test_classify_kind_detects_business_card_from_image_with_contact_info():
    ocr_text = "株式会社サンプル\n営業部 田中太郎\nTEL: 03-1234-5678\nEmail: tanaka@example.com"
    assert svc.classify_kind(file_name="IMG_0001.jpg", extracted_text=ocr_text, content_type="image/jpeg") == "business_card"


def test_classify_kind_image_without_contact_info_is_not_business_card():
    assert svc.classify_kind(file_name="photo.jpg", extracted_text="", content_type="image/jpeg") == "other"


def test_classify_kind_non_image_with_contact_info_is_not_business_card():
    # A regular document can contain a phone number too — only images
    # should ever be considered for the business_card heuristic.
    text = "会議の議事録です。連絡先は 03-1234-5678 までお願いします。"
    assert svc.classify_kind(file_name="minutes.pdf", extracted_text=text, content_type="application/pdf") == "other"


def test_extract_text_docx_round_trip():
    import docx

    buffer = io.BytesIO()
    document = docx.Document()
    document.add_paragraph("案件名: Javaエンジニア募集")
    document.add_paragraph("単価: 70万円")
    document.save(buffer)

    text = svc.extract_text(
        file_name="案件票.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        data=buffer.getvalue(),
    )
    assert "Javaエンジニア募集" in text
    assert "70万円" in text


def test_extract_text_xlsx_round_trip():
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["氏名", "スキル", "単価"])
    sheet.append(["A氏", "Java, Spring", "65万円"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    text = svc.extract_text(
        file_name="skill_sheet.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        data=buffer.getvalue(),
    )
    assert "Java, Spring" in text
    assert "65万円" in text


def test_extract_text_invalid_pdf_bytes_returns_empty_string_not_raise():
    text = svc.extract_text(file_name="broken.pdf", content_type="application/pdf", data=b"not a real pdf")
    assert text == ""


def test_extract_text_plain_text_uses_utf8_decode():
    text = svc.extract_text(file_name="notes.txt", content_type="text/plain", data="メモ書き".encode("utf-8"))
    assert text == "メモ書き"


def test_extract_text_zip_recurses_into_members():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("メモ.txt", "Java案件の候補者リストです")
        archive.writestr("請求書.txt", "請求書 合計 300,000円")

    text = svc.extract_text(file_name="資料.zip", content_type="application/zip", data=buffer.getvalue())
    assert "Java案件の候補者リストです" in text
    assert "300,000円" in text
    assert "=== メモ.txt ===" in text


def test_extract_text_zip_resolved_by_extension_when_content_type_missing():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("notes.txt", "hello from inside the zip")

    text = svc.extract_text(file_name="archive.zip", content_type="application/octet-stream", data=buffer.getvalue())
    assert "hello from inside the zip" in text


def test_extract_text_invalid_zip_bytes_returns_empty_string_not_raise():
    text = svc.extract_text(file_name="broken.zip", content_type="application/zip", data=b"not a real zip")
    assert text == ""


def test_extract_text_zip_skips_members_beyond_max_count(monkeypatch):
    monkeypatch.setattr(svc, "_MAX_ZIP_MEMBERS", 2)

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for i in range(5):
            archive.writestr(f"file_{i}.txt", f"content {i}")

    text = svc.extract_text(file_name="many.zip", content_type="application/zip", data=buffer.getvalue())
    included = sum(1 for i in range(5) if f"content {i}" in text)
    assert included == 2


def test_analyze_classifies_kind_from_zip_member_content():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("doc.txt", "本書は基本契約書として締結する")

    result = svc.analyze(file_name="書類.zip", content_type="application/zip", data=buffer.getvalue())
    assert result.kind == "contract"


def test_analyze_returns_kind_and_text_together():
    import docx

    buffer = io.BytesIO()
    document = docx.Document()
    document.add_paragraph("請求書 合計金額 500,000円")
    document.save(buffer)

    result = svc.analyze(
        file_name="invoice.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        data=buffer.getvalue(),
    )
    assert result.kind == "invoice"
    assert "500,000円" in result.extracted_text
