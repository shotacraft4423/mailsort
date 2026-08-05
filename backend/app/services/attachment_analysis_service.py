"""添付解析: extracts text from PDF/Excel/Word attachments and classifies
what kind of document they are (skill sheet / project brief / invoice /
contract / other) so the UI can show a dedicated icon and the classification
service can feed the extracted text into its prompt context.

Image OCR (名刺OCR含む) needs a `pytesseract` + Tesseract binary install
that isn't guaranteed to be present on every machine; `extract_text` degrades
gracefully (returns "") when the OCR stack is missing rather than raising,
so attachment ingestion never blocks on it. Zip-file recursion is a Phase 2
item (would unzip into a temp dir and re-run this same function per member).
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass

AttachmentKind = str  # skill_sheet | project_brief | invoice | contract | resume | other

_KIND_KEYWORDS: dict[str, list[str]] = {
    "skill_sheet": ["スキルシート", "skill sheet", "経歴書", "職務経歴"],
    "invoice": ["請求書", "invoice", "御請求", "お振込み先"],
    "contract": ["契約書", "基本契約", "発注書", "個別契約", "contract"],
    "project_brief": ["案件票", "案件概要", "募集要項"],
    "resume": ["履歴書", "resume", "cv"],
}


@dataclass
class AttachmentAnalysis:
    kind: AttachmentKind
    extracted_text: str


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        return ""
    try:
        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception:  # noqa: BLE001 - a malformed/encrypted PDF must not break ingestion
        return ""


def _extract_docx(data: bytes) -> str:
    try:
        import docx
    except ImportError:
        return ""
    try:
        document = docx.Document(io.BytesIO(data))
        return "\n".join(p.text for p in document.paragraphs)
    except Exception:  # noqa: BLE001
        return ""


def _extract_xlsx(data: bytes) -> str:
    try:
        import openpyxl
    except ImportError:
        return ""
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lines = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    lines.append("\t".join(cells))
        return "\n".join(lines)
    except Exception:  # noqa: BLE001
        return ""


def _extract_image_ocr(data: bytes) -> str:
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        return ""
    try:
        return pytesseract.image_to_string(Image.open(io.BytesIO(data)), lang="jpn+eng")
    except Exception:  # noqa: BLE001 - missing tesseract binary, corrupt image, etc.
        return ""


_EXTRACTORS = {
    "application/pdf": _extract_pdf,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": _extract_docx,
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": _extract_xlsx,
    "image/png": _extract_image_ocr,
    "image/jpeg": _extract_image_ocr,
}

_EXTENSION_CONTENT_TYPE = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
}


def _resolve_content_type(file_name: str, content_type: str) -> str:
    if content_type in _EXTRACTORS:
        return content_type
    for ext, mapped in _EXTENSION_CONTENT_TYPE.items():
        if file_name.lower().endswith(ext):
            return mapped
    return content_type


def extract_text(*, file_name: str, content_type: str, data: bytes) -> str:
    resolved = _resolve_content_type(file_name, content_type)
    extractor = _EXTRACTORS.get(resolved)
    if extractor is None:
        if resolved.startswith("text/"):
            return data.decode("utf-8", errors="replace")
        return ""
    return extractor(data).strip()


def classify_kind(*, file_name: str, extracted_text: str) -> AttachmentKind:
    haystack = f"{file_name}\n{extracted_text[:2000]}".lower()
    for kind, keywords in _KIND_KEYWORDS.items():
        if any(re.search(re.escape(kw.lower()), haystack) for kw in keywords):
            return kind
    return "other"


def analyze(*, file_name: str, content_type: str, data: bytes) -> AttachmentAnalysis:
    text = extract_text(file_name=file_name, content_type=content_type, data=data)
    kind = classify_kind(file_name=file_name, extracted_text=text)
    return AttachmentAnalysis(kind=kind, extracted_text=text)
